"""edit_lib.py — 개조식 변환 · 글머리 「- 」 · 나열 묶음 줄바꿈 · 들여쓰기 · [[ ]] 빨강(#FF0000) 유지 글 넣기 (mgc-ppt-template v5.4, references/text-fit.md 「묶어 나누지 않는 말」)"""
import copy, re
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn
import pptx_edit as PE

BULLET = '- '          # 정보 단위 글머리: references/writing.md 「기호」
RED = 'FF0000'; INK = '2E241D'   # 강조 빨강 단일 값: references/writing.md 「강조」(이전 값은 먼저 PE.fix_colors로 정리)
SLACK = PE.SLACK       # 줄바꿈 폭 보정(pptx_edit와 같은 값): references/text-fit.md 「폭 추정 · 보정 값」

_em0 = PE.em
def em(ch, bold):
    if ch == '•': return 0.62
    if ch == '\u00a0': return 0.30
    return _em0(ch, bold)
PE.em = em
def wpt(s, size, bold): return sum(em(c, bold) * size for c in re.sub(r'\[\[|\]\]', '', s))

# ── 줄바꿈: 「A · B · C」 나열은 한 덩어리로 묶어 줄 끝 · 줄 머리에 「·」가 남지 않게 한다
def _units(par):
    """띄어쓰기 단위를 묶는다: 짧은 괄호 안(괄호 시작부터 24자 이하) · 「A · B」 나열 · 「[확인필요(…)]」는 한 덩어리"""
    toks = par.split(' '); units = []; glue = False; depth = 0; plen = 0
    for t in toks:
        inner = depth > 0 and plen <= 24
        if units and (glue or inner or t == '·'):
            units[-1] += ' ' + t
        else:
            units.append(t)
        glue = (t == '·')
        before = depth
        depth += t.count('(') + t.count('[') - t.count(')') - t.count(']')
        depth = max(depth, 0)
        if depth > 0:
            if before == 0:
                k = max(t.rfind('('), t.rfind('['))
                plen = len(t) - k
            else: plen += len(t) + 1
            if units[-1].startswith('[확인필요') or '[확인필요' in t: plen = 0
        else: plen = 0
    return units

GLUE = ['CJ ONS', '메가MGC커피 서버', '핏콜라보 서버', 'App Store', 'Play Console', '확인 필요', '처리 흐름', '연동 구조', '정책 기준',
        '카드관리 페이지', '공통 웹 컴포넌트', '브릿지 페이지', '인증 모듈', '결제 화면', '고객 화면', '주문 결과']
TAG_RE = re.compile(r'\[[^\[\]]*?(?:확인|검토) 필요\]')   # 서술형 확인 태그: 한 덩어리
def wrap(par, w_in, size, bold=False):
    glued = par
    for g in GLUE: glued = glued.replace(g, g.replace(' ', '\u00a0'))
    plain = glued.replace('확인\u00a0필요', '확인 필요').replace('검토\u00a0필요', '검토 필요')   # GLUE가 묶은 「확인 필요」를 풀어 태그 전체를 찾는다
    tagged = TAG_RE.sub(lambda m: m.group(0).replace(' ', '\u00a0'), plain)
    avail = w_in * 72 * SLACK
    out = _wrap(tagged, w_in, size, bold)
    if any(wpt(o, size, bold) > avail for o in out):   # 태그가 한 줄보다 길면 태그 안에서만 나눈다
        out = _wrap(glued, w_in, size, bold)
    return [x.replace('\u00a0', ' ') for x in out]
def _wrap(par, w_in, size, bold=False):
    """한 정보(한 문단)를 칸 폭에 맞춰 줄로 나눈다. 반환: [첫 줄, 이어지는 줄...]
    나열만으로 된 정보(예: 「정책(12~23) · 화면(25~31) · 예외(59~64)」)는 줄 경계의 「·」를 빼고,
    문장 속 나열이 한 줄을 넘으면 「·」를 줄 끝에 남긴다"""
    avail = w_in * 72 * SLACK
    lead = BULLET if par.startswith(BULLET) else ''
    body = par[len(lead):]
    ind = wpt(lead, size, bold)
    us = _units(body)
    pure = len(us) == 1 and ' · ' in body
    out, cur = [], ''
    def fits(s): return wpt(s, size, bold) + ind <= avail
    us2 = []
    for u in us:   # 한 덩어리가 한 줄보다 길면(나열 아님) 띄어쓰기 단위로 푼다
        if not fits(u) and ' · ' not in u: us2.extend(u.split(' '))
        else: us2.append(u)
    us = us2
    for u in us:
        cand = u if not cur else cur + ' ' + u
        if cur and not fits(cand):
            out.append(cur); cur = u
        else:
            cur = cand
        if not fits(cur) and ' · ' in cur:
            parts = cur.split(' · '); line = parts[0]
            for k, part in enumerate(parts[1:], 1):
                last = k == len(parts) - 1
                trial = line + ' · ' + part + ('' if (last or pure) else ' ·')
                if fits(trial): line = line + ' · ' + part
                else:
                    out.append(line if pure else line + ' ·'); line = part
            cur = line
    out.append(cur)
    out = [o for o in out if o]
    out[0] = lead + out[0]
    return out

# ── 개조식 끝맺음 변환(표 칸 · 설명 · 각주 전용)
JONG_B, JONG_M, JONG_N = 17, 16, 4
def _jong(ch, j):
    c = ord(ch) - 0xAC00
    return chr(0xAC00 + c - c % 28 + j)
def _has_jong(ch, j=None):
    if not ('가' <= ch <= '힣'): return False
    k = (ord(ch) - 0xAC00) % 28
    return k != 0 if j is None else k == j
VERB_I = {'붙입', '보입', '쓰입', '줄입', '놓입', '높입', '들입', '쌓입', '먹입', '이깁'}
SPECIAL = {'엽니다': '열어 줌', '만듭니다': '만듦', '겁니다': '걺', '압니다': '앎', '아닙니다': '아님'}

def _end(word):
    """끝맺음 한 어절을 개조식으로. 바뀌지 않으면 None"""
    for k, v in SPECIAL.items():
        if word.endswith(k): return word[:-len(k)] + v
    if word.endswith('습니다'): return word[:-3] + '음'
    if word.endswith('입니다'):
        if len(word) >= 4 and word[-4:-2] in VERB_I: return word[:-3] + _jong(word[-3], JONG_M)
        return word[:-3]                  # 「X입니다」 → 「X」(숫자 · 괄호 뒤 「입니다」는 빈 문자열)
    if word.endswith('니다') and len(word) >= 3 and _has_jong(word[-3], JONG_B):
        return word[:-3] + _jong(word[-3], JONG_M)
    return None

END_RE = re.compile(r'([가-힣]+니다)\.?(?=(\]\])?(\s*\[확인필요[^\]]*\])?\s*(\([^)]*\))?\s*$)')
def gaejo(line):
    """줄 끝 서술어(~합니다 · ~습니다 · ~입니다)를 개조식으로 바꾸고 끝 온점을 뺀다. 문장 중간 「~합니다,」도 바꾼다"""
    def rep(m):
        new = _end(m.group(1))
        return new if new is not None else m.group(1)
    s = re.sub(r'([가-힣]+니다)\.(?=\s|$)', lambda m: (_end(m.group(1)) if _end(m.group(1)) is not None else m.group(1)), line)
    s = END_RE.sub(lambda m: rep(m), s)
    s = re.sub(r'([가-힣]+니다)(?=, )', lambda m: (_end(m.group(1)) if _end(m.group(1)) is not None else m.group(1)), s)
    s = re.sub(r'(\])\.$', r'\1', s.rstrip())
    return s

# ── 칸 · 글상자의 논리 정보(문단) 복원: 자동 줄바꿈으로 나뉜 줄을 다시 붙인다
def items_of(tf):
    lines = [pa.text for pa in tf.paragraphs]
    out = []
    for ln in lines:
        raw = ln
        s = ln.strip()
        if not s: continue
        is_head = s.startswith(('- ', '• ', '※'))
        if s.startswith('• '): s = BULLET + s[2:]
        cont = (raw.startswith('  ') or not is_head)
        if out and cont and not out[-1].endswith('.') and not out[-1].endswith(']]'):
            out[-1] = out[-1] + ('' if out[-1].endswith(('(', '/')) else ' ') + s
        else:
            out.append(s)
    return out

def red_text(tf):
    """빨강 강조 런을 [[ ]]로 표시한 문단 목록(강조 유지용)"""
    res = []
    for pa in tf.paragraphs:
        s = ''
        for r in pa.runs:
            rp = r._r.find(qn('a:rPr')); c = rp.find('.//' + qn('a:srgbClr')) if rp is not None else None
            s += f'[[{r.text}]]' if (c is not None and c.get('val', '').upper() == RED) else r.text
        res.append(s.replace(']][[', ''))
    return res

# ── 글 넣기: [[ ]] 빨강, 글머리 이어지는 줄은 문단 들여쓰기(marL)로 맞춘다
def set_items(tf, items, w_in, size=None, bold=None, plain_color=INK):
    sz, b = PE.run_props(tf); size = size or sz; bold = b if bold is None else bold
    body = tf._txBody; ps = body.findall(qn('a:p')); p0 = ps[0]
    allr = [r for p in ps for r in p.findall(qn('a:r'))]
    def isred(r):
        rp = r.find(qn('a:rPr')); c = rp.find('.//' + qn('a:srgbClr')) if rp is not None else None
        return c is not None and c.get('val', '').upper() == RED
    nonred = [r for r in allr if not isred(r)]
    tmpl = copy.deepcopy(nonred[0] if nonred else allr[0])
    if not nonred:   # 모든 런이 빨강이었던 칸: 본문색으로 되돌린다
        rp = tmpl.find(qn('a:rPr'))
        for f in rp.findall(qn('a:solidFill')): rp.remove(f)
        sf = rp.makeelement(qn('a:solidFill'), {}); sf.append(sf.makeelement(qn('a:srgbClr'), {'val': plain_color})); rp.insert(0, sf)
        rp.set('b', '1' if bold else '0')
    for p in ps: body.remove(p)
    n = 0
    for it in items:
        if it.startswith('[[- '): it = '- [[' + it[4:]
        lines = wrap(it, w_in, size, bold) if w_in else [it]
        fixed, carry = [], False
        for ln in lines:
            if carry: ln = '[[' + ln
            o, c = ln.count('[['), ln.count(']]')
            carry = o > c
            if carry: ln = ln + ']]'
            fixed.append(ln.replace('[[]]', ''))
        lines = fixed
        for li, ln in enumerate(lines):
            np_ = copy.deepcopy(p0)
            for r in np_.findall(qn('a:r')) + np_.findall(qn('a:br')) + np_.findall(qn('a:fld')): np_.remove(r)
            ppr = np_.find(qn('a:pPr'))
            if ppr is None: ppr = np_.makeelement(qn('a:pPr'), {}); np_.insert(0, ppr)
            ppr.set('eaLnBrk', '0')   # 단어 끊김 방지
            if li > 0 and it.startswith(BULLET):
                ppr.set('marL', str(int(wpt(BULLET, size, bold) * 12700))); ppr.set('indent', '0')
            else:
                for k in ('marL', 'indent'):
                    if k in ppr.attrib: del ppr.attrib[k]
            end = np_.find(qn('a:endParaRPr'))
            red_open = False
            for part in re.split(r'(\[\[.*?\]\](?!\]))', ln):   # 「[[[태그]]]」처럼 괄호까지 빨강인 경우 포함
                if not part: continue
                red = part.startswith('[[') and part.endswith(']]')
                nr = copy.deepcopy(tmpl); nr.find(qn('a:t')).text = part[2:-2] if red else part
                if red:
                    rp = nr.find(qn('a:rPr')); rp.set('b', '1')
                    for f in rp.findall(qn('a:solidFill')): rp.remove(f)
                    sf = rp.makeelement(qn('a:solidFill'), {}); sf.append(sf.makeelement(qn('a:srgbClr'), {'val': RED})); rp.insert(0, sf)
                (end.addprevious(nr) if end is not None else np_.append(nr))
            body.append(np_); n += 1
    return n

def cell_w(table, ci): return table.columns[ci].width / 914400 - 0.24
def row_h(table, ri, minh=0.36):
    mx = 0
    for c in table.rows[ri].cells:
        s, _ = PE.run_props(c.text_frame); mx = max(mx, len(c.text_frame.paragraphs) * s * 1.3 / 72)
    return max(minh, mx + 0.14)
