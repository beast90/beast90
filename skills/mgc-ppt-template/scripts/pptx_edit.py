"""pptx_edit.py — 사용자가 다듬은 pptx를 python-pptx로 직접 고칠 때 쓰는 도구 (mgc-ppt-template v5.4, references/edit-user-pptx.md)

from pptx_edit import *
- 줄나눔: wrap(text, w_in, size, bold) — 칸 폭 기준 띄어쓰기 단위, '- ' 항목 이어지는 줄은 HANG 표시(쓸 때 문단 들여쓰기 marL로 바꿈, 공백 들여쓰기 없음)
- 글: set_tf(tf, text, w_in, nowrap) — 빨강이 아닌 첫 런 서식을 원본으로 문단 재구성, [[문구]]는 빨강 굵게
- 표: set_cell · fit_rows · del_row · add_col · del_col · tbl · bottom
- 도형: textbox · box · arrow (맑은 고딕, 정보 도형만)
- 장: duplicate_slide · renumber(p, inserted_at) · add_source_line · replace_picture
- 표 정리: balance_cols(table) — 열 폭 합은 두고 문장 열 폭을 다시 나눠 한 줄에 들어가게, rewrap_table · fit_rows로 줄 · 행 높이 다시 계산
- 마무리: fix_hang(p) 공백 들여쓰기 → 문단 들여쓰기, no_word_split(p) 모든 문단 eaLnBrk="0"(한글 단어 중간 줄바꿈 금지). 명령: python3 pptx_edit.py finalize 입력.pptx 출력.pptx
- 색: fix_colors(p) — 토큰 밖 이전 값(회색 7A6A5E, 빨강 F53148 · D32F2F · E03C31 · C00000)을 토큰 값으로 일괄 교체
"""
import copy, re, sys
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ── 폭 추정 (fit_text.js와 같은 표). 줄바꿈 폭 보정 SLACK: references/text-fit.md 「폭 추정 · 보정 값」
def em(ch, bold):
    c = ord(ch); b = 1.06 if bold else 1
    if (0xAC00 <= c <= 0xD7A3) or (0x1100 <= c <= 0x11FF) or (0x3130 <= c <= 0x318F) or (0x4E00 <= c <= 0x9FFF) or (0x3000 <= c <= 0x303F) or (0xFF00 <= c <= 0xFFEF):
        return 0.98 if bold else 0.96
    if ch == ' ': return 0.30
    if ch.isdigit(): return 0.56 * b
    if 'A' <= ch <= 'Z': return 0.64 * b
    if 'a' <= ch <= 'z': return 0.52 * b
    if ch in ".,:;'\"`|!‘’“”·": return 0.30
    if ch in '()[]{}': return 0.34
    if ch in '→←↑↓↔—@': return 1.0
    if ch in '~+=<>#': return 0.62
    if ch == '%': return 0.85
    if ch in '/\\': return 0.40
    if ch == '-': return 0.38
    return 0.92
def wpt(s, size, bold): return sum(em(c, bold) * size for c in s)
HANG = '\x01'   # 「- 」 항목의 이어지는 줄 표시: 문단 들여쓰기(marL)로 그림
SLACK = 0.95   # 사용자 수정 표(앱카드 연동 통합본 v0.26 목차) 실측: 추정 폭 = 실제의 0.95~0.98배
def wrap(text, w_in, size, bold=False, slack=SLACK):
    avail = w_in * 72 * slack; out = []
    hang_w = wpt('- ', size, bold)
    for par in str(text).split('\n'):
        hy = par.startswith('- '); cur = ''
        for tok in par.split(' '):
            cand = tok if cur == '' else cur + ' ' + tok
            w = wpt(re.sub(r'\[\[|\]\]', '', cand.lstrip(HANG)), size, bold) + (hang_w if cand.startswith(HANG) else 0)
            if cur != '' and w > avail:
                out.append(cur); cur = (HANG if hy else '') + tok
            else: cur = cand
        out.append(cur)
    return out
def _ppr(p_el):
    pPr = p_el.find(qn('a:pPr'))
    if pPr is None: pPr = p_el.makeelement(qn('a:pPr'), {}); p_el.insert(0, pPr)
    return pPr
def _hang(p_el, ln, size):
    """「- 」 줄은 내어쓰기, 이어지는 줄은 같은 폭 들여쓰기. 단어 끊김 방지(eaLnBrk=0)도 같이 넣는다. 반환: HANG을 뗀 글"""
    pPr = _ppr(p_el); pPr.set('eaLnBrk', '0')
    if ln.startswith('- ') or ln.startswith(HANG):
        if pPr is None: pPr = p_el.makeelement(qn('a:pPr'), {}); p_el.insert(0, pPr)
        m = str(int(round(wpt('- ', size, False) * 12700)))
        pPr.set('marL', m); pPr.set('indent', '-' + m if ln.startswith('- ') else '0')
    elif pPr is not None:
        for k in ('marL', 'indent'): pPr.attrib.pop(k, None)
    return ln.lstrip(HANG)

RED = {'F53148', 'D32F2F', 'E03C31', 'C00000', 'FF0000'}   # 이전 값도 빨강 런으로 인식(쓸 때는 RED_HEX만)
RED_HEX = 'FF0000'
def _isred(r):
    rp = r.find(qn('a:rPr'))
    if rp is None: return False
    c = rp.find('.//' + qn('a:srgbClr'))
    return c is not None and c.get('val', '').upper() in RED
def run_props(tf):
    for p in tf.paragraphs:
        for r in p.runs:
            rp = r._r.find(qn('a:rPr'))
            return (int(rp.get('sz')) / 100 if rp is not None and rp.get('sz') else 12), (rp is not None and rp.get('b') == '1')
    return 12, False
def set_tf(tf, text, w_in=None, size=None, bold=None, nowrap=False):
    sz, b = run_props(tf); size = size or sz; bold = b if bold is None else bold
    lines = str(text).split('\n') if (nowrap or w_in is None) else wrap(text, w_in, size, bold)
    body = tf._txBody; ps = body.findall(qn('a:p')); p0 = ps[0]
    allr = [r for p in ps for r in p.findall(qn('a:r'))]
    nonred = [r for r in allr if not _isred(r)]
    tmpl = copy.deepcopy(nonred[0] if nonred else allr[0])
    for p in ps: body.remove(p)
    for ln in lines:
        np_ = copy.deepcopy(p0)
        for r in np_.findall(qn('a:r')) + np_.findall(qn('a:br')) + np_.findall(qn('a:fld')): np_.remove(r)
        end = np_.find(qn('a:endParaRPr'))
        ln = _hang(np_, ln, size)
        for part in re.split(r'(\[\[.*?\]\])', ln):
            if not part: continue
            red = part.startswith('[[')
            nr = copy.deepcopy(tmpl); nr.find(qn('a:t')).text = part[2:-2] if red else part
            if red:
                rp = nr.find(qn('a:rPr')); rp.set('b', '1')
                for f in rp.findall(qn('a:solidFill')): rp.remove(f)
                sf = rp.makeelement(qn('a:solidFill'), {}); sf.append(sf.makeelement(qn('a:srgbClr'), {'val': RED_HEX})); rp.insert(0, sf)
            (end.addprevious(nr) if end is not None else np_.append(nr))
        body.append(np_)
    return len(lines)

# ── 표
def tbl(slide, idx=0): return [sh for sh in slide.shapes if sh.has_table][idx]
def find(slide, name): return [sh for sh in slide.shapes if sh.name == name]
def bottom(sh):
    return (sh.top + sum(r.height for r in sh.table.rows)) / 914400 if sh.has_table else (sh.top + sh.height) / 914400
def set_cell(table, ri, ci, text):
    return set_tf(table.cell(ri, ci).text_frame, text, table.columns[ci].width / 914400 - 0.24)
def marked(pa):
    """문단 글. 빨강 런은 [[ ]]로 표시(다시 넣을 때 강조 유지)"""
    return ''.join(f'[[{r.text}]]' if _isred(r._r) else r.text for r in pa.runs).replace(']][[', '')
def merged_text(cell):
    """자동으로 나뉜 줄을 다시 한 정보로 합친다('- ' 항목과 온점 끝 줄은 유지, 빨강 강조 유지)"""
    out = []
    for ln in [marked(pa) for pa in cell.text_frame.paragraphs]:
        s = ln.strip()
        if out and not s.startswith('- ') and not out[-1].rstrip().endswith('.'): out[-1] = out[-1].rstrip() + ' ' + s
        else: out.append(s)
    return '\n'.join(out)
def rewrap_table(table):
    for ri in range(len(table.rows)):
        for ci in range(len(table.columns)): set_cell(table, ri, ci, merged_text(table.cell(ri, ci)))
def balance_cols(table, cols=None, slack=SLACK):
    """열 폭 합은 그대로 두고 문장 열의 폭을 다시 나눠 정보(「- 」 항목)마다 한 줄에 들어가게 한다.
    cols: 다시 나눌 열(기본: 첫 열을 뺀 열 중 평균 10자 이상이거나 「- 」로 시작하는 열). 반환: 새 폭 목록 또는 None(합이 모자라 그대로 둠)"""
    nr, nc = len(table.rows), len(table.columns); start = 1 if nr > 1 else 0
    def items(ci):
        out = []
        for ri in range(start, nr):
            c = table.cell(ri, ci); size, bold = run_props(c.text_frame)
            out += [(re.sub(r'\[\[|\]\]', '', t), size, bold) for t in merged_text(c).split('\n') if t.strip()]
        return out
    if cols is None:
        cols = []
        for ci in range(1, nc):
            its = items(ci)
            if its and (any(t.startswith('- ') for t, _, _ in its) or sum(len(t) for t, _, _ in its) / len(its) >= 10): cols.append(ci)
    if len(cols) < 2: return None
    need = {ci: max((wpt(t, sz, b) for t, sz, b in items(ci)), default=0) / (72 * slack) + 0.24 + 0.01 for ci in cols}
    total = sum(table.columns[ci].width for ci in cols) / 914400
    if sum(need.values()) > total + 1e-6: return None
    extra = (total - sum(need.values())) / len(cols)
    ws = {ci: need[ci] + extra for ci in cols}
    ws[cols[-1]] = total - sum(ws[c] for c in cols[:-1])   # 합 오차는 마지막 열에서 맞춤
    for ci in cols: table.columns[ci].width = Inches(ws[ci])
    return [round(table.columns[ci].width / 914400, 3) for ci in range(nc)]
def fit_rows(table, rows=None, minh=0.36):
    for ri, row in enumerate(table.rows):
        if rows is not None and ri not in rows: continue
        mx = 0
        for c in row.cells:
            s, _ = run_props(c.text_frame); mx = max(mx, len(c.text_frame.paragraphs) * s * 1.3 / 72)
        row.height = Inches(max(minh, mx + 0.14))
def set_size(table, size):
    for e in list(table._tbl.iter(qn('a:rPr'))) + list(table._tbl.iter(qn('a:endParaRPr'))): e.set('sz', str(int(size * 100)))
def del_row(table, ri): table._tbl.remove(table._tbl.tr_lst[ri])
def add_row(table): table._tbl.append(copy.deepcopy(table._tbl.tr_lst[-1]))
def add_col(table, src=1, at=1):
    grid = table._tbl.find(qn('a:tblGrid')); grid.insert(at, copy.deepcopy(grid.findall(qn('a:gridCol'))[src]))
    for tr in table._tbl.tr_lst:
        tcs = tr.findall(qn('a:tc')); tr.insert(list(tr).index(tcs[at]), copy.deepcopy(tcs[src]))
def del_col(table, ci):
    grid = table._tbl.find(qn('a:tblGrid')); grid.remove(grid.findall(qn('a:gridCol'))[ci])
    for tr in table._tbl.tr_lst: tr.remove(tr.findall(qn('a:tc'))[ci])

# ── 도형 (정보 도형만, 1-1)
INK = RGBColor(0x2E, 0x24, 0x1D); MUTED = RGBColor(0x7A, 0x6A, 0x5D); BROWN = RGBColor(0x3B, 0x2A, 0x20)
BEIGE = RGBColor(0xF3, 0xEE, 0xE8); LINE = RGBColor(0xC9, 0xBB, 0xAA); FC_RED = RGBColor(0xFF, 0x00, 0x00)
FC_PINK = RGBColor(0xFD, 0xEC, 0xEC); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
def _font(r, size, bold, color):
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    rp = r._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        e = rp.find(qn(tag))
        if e is None: e = rp.makeelement(qn(tag), {}); rp.append(e)
        e.set('typeface', 'Malgun Gothic')
def textbox(s, x, y, w, h, lines, size=11, color=INK, bold=False, align=PP_ALIGN.LEFT, name='본문'):
    out = []
    for ln in lines: out += wrap(ln, w, size, bold)
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tb.name = name
    tf = tb.text_frame; tf.word_wrap = True
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    for i, ln in enumerate(out):
        pa = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); pa.alignment = align
        ln = _hang(pa._p, ln, size)
        _font(pa.add_run(), size, bold, color); pa.runs[0].text = ln
    tb.height = Inches(len(out) * size * 1.3 / 72 + 0.04)
    return tb
def box(s, x, y, w, h, lines, fill=BEIGE, colors=None, sizes=None, bolds=None, red_border=False, name='도식 상자'):
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); b.name = name
    b.fill.solid(); b.fill.fore_color.rgb = fill
    b.line.color.rgb = FC_RED if red_border else LINE; b.line.width = Pt(1.25 if red_border else 0.75)
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06); tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, ln in enumerate(lines):
        pa = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); pa.alignment = PP_ALIGN.CENTER
        r = pa.add_run(); r.text = ln
        _font(r, (sizes or [11.5] + [10] * 9)[i], (bolds or [True] + [False] * 9)[i], (colors or [INK] + [MUTED] * 9)[i])
    return b
def arrow(s, x1, y1, x2, y2, color=MUTED, name='도식 화살표'):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); c.name = name
    c.line.color.rgb = color; c.line.width = Pt(1.25)
    ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    return c

# ── 장
def duplicate_slide(p, src_index, insert_at):
    """src_index 장을 복제해 insert_at(0부터) 위치에 넣는다. 노트는 노트 틀이 있을 때만 복사"""
    src = p.slides[src_index]; new = p.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    for el in src.shapes._spTree.iterchildren():
        if el.tag.split('}')[1] in ('sp', 'graphicFrame', 'cxnSp', 'grpSp', 'pic'): new.shapes._spTree.append(copy.deepcopy(el))
    ntf = new.notes_slide.notes_text_frame
    if ntf is not None and src.has_notes_slide: ntf.text = src.notes_slide.notes_text_frame.text
    ids = p.slides._sldIdLst; last = list(ids)[-1]; ids.remove(last); ids.insert(insert_at, last)
    return p.slides[insert_at]
def _runs(s):
    for sh in s.shapes:
        if sh.has_text_frame:
            for pa in sh.text_frame.paragraphs:
                for r in pa.runs: yield sh, r
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    for pa in c.text_frame.paragraphs:
                        for r in pa.runs: yield sh, r
def renumber(p, new_no):
    """new_no(1부터) 장을 새로 넣은 뒤: 쪽 번호 n / N, 간지 목록 번호, 본문 'N장' · 'N~M장' · '(N~M)' 참조를 다시 맞춘다"""
    N = len(p.slides); PG = re.compile(r'^\d+ / \d+$')
    bump = lambda n: n + 1 if n >= new_no else n
    RNG = re.compile(r'(\d{1,2})~(\d{1,2})(?=\)|장)'); ONE = re.compile(r'(?<![\d.~])(\d{1,2})(?=장)')
    for i, s in enumerate(p.slides, 1):
        for sh, r in _runs(s):
            t = r.text
            if PG.match(t.strip()): r.text = f'{i} / {N}'; continue
            if sh.name == '간지 목록':
                m = re.match(r'^(\d+)(\s+)(.*)$', t)
                if m and i != new_no: r.text = f'{bump(int(m.group(1)))}{m.group(2)}{m.group(3)}'
                continue
            t2 = RNG.sub(lambda m: f'{bump(int(m.group(1)))}~{int(m.group(2)) + 1 if int(m.group(2)) >= new_no - 1 else int(m.group(2))}', t)
            t2 = ONE.sub(lambda m: str(bump(int(m.group(1)))), t2)
            if t2 != t: r.text = t2
def add_source_line(p, text, skip=(), pattern=r'[①②③④⑤⑥⑦⑧⑨]'):
    """번호를 쓰는 장 하단(y 7.02)에 출처 1줄. skip: 표지 · 간지 · 번호 원본 장 번호"""
    pat = re.compile(pattern)
    for i, s in enumerate(p.slides, 1):
        if i in skip or '간지 목록' in [x.name for x in s.shapes]: continue
        if pat.search(' '.join(r.text for _, r in _runs(s))) and not find(s, '출처'):
            textbox(s, 0.6, 7.02, 10.9, 0.26, [text], size=9, color=MUTED, name='출처')
def replace_picture(slide, png_path):
    """그림 위치 · 크기는 그대로 두고 이미지만 교체(같은 비율로 자른 이미지일 것)"""
    pic = [x for x in slide.shapes if x.shape_type == 13][0]
    pic.part.related_part(pic._element.blipFill.blip.rEmbed)._blob = open(png_path, 'rb').read()
    return pic


# ── 색 정리: 토큰 밖 옛 값을 토큰 값으로 (글자 · 채우기 · 선 모두)
def no_word_split(p):
    """모든 글상자 · 표 칸 문단에 eaLnBrk="0": 줄이 넘쳐도 한글 단어 가운데가 아니라 띄어쓰기에서 줄을 바꾼다. 반환: 바꾼 문단 수"""
    n = 0
    for s in p.slides:
        for p_el in s._element.iter(qn('a:p')):
            if p_el.find(qn('a:r')) is None: continue
            pPr = _ppr(p_el)
            if pPr.get('eaLnBrk') != '0': pPr.set('eaLnBrk', '0'); n += 1
    return n

def fix_hang(p):
    """「- 」 문단 뒤 공백 두 칸으로 시작하는 이어지는 줄(생성 코드 · 이전 판 표시)을 문단 들여쓰기로 바꾼다. 반환: 바꾼 문단 수"""
    n = 0
    for s in p.slides:
        for body in s._element.iter(qn('a:txBody')):
            prev_dash = False
            for p_el in body.findall(qn('a:p')):
                ts = [t for t in p_el.iter(qn('a:t'))]; txt = ''.join(t.text or '' for t in ts)
                rp = p_el.find('.//' + qn('a:rPr')); size = int(rp.get('sz')) / 100 if rp is not None and rp.get('sz') else 12
                if txt.startswith('- '):
                    prev_dash = True; pPr = _ppr(p_el); m = str(int(round(wpt('- ', size, False) * 12700)))
                    pPr.set('marL', m); pPr.set('indent', '-' + m)
                elif prev_dash and txt.startswith('  ') and ts:
                    k = len(txt) - len(txt.lstrip(' '))
                    for t in ts:   # 앞 공백만 지운다(여러 런에 걸쳐도)
                        if k <= 0: break
                        cut = min(k, len(t.text or '') - len((t.text or '').lstrip(' '))); t.text = (t.text or '')[cut:]; k -= cut
                        if (t.text or '').strip(): break
                    pPr = _ppr(p_el); pPr.set('marL', str(int(round(wpt('- ', size, False) * 12700)))); pPr.set('indent', '0'); n += 1
                elif txt.strip(): prev_dash = False
    return n

LEGACY = {'7A6A5E': '7A6A5D', 'F53148': RED_HEX, 'D32F2F': RED_HEX, 'E03C31': RED_HEX, 'C00000': RED_HEX}   # 이전 값 → 토큰
def fix_colors(p, mapping=LEGACY):
    n = 0
    for s in p.slides:
        for e in s._element.iter(qn('a:srgbClr')):
            v = e.get('val', '').upper()
            if v in mapping: e.set('val', mapping[v]); n += 1
    return n

if __name__ == '__main__' and len(sys.argv) == 4 and sys.argv[1] == 'finalize':
    from pptx import Presentation
    _p = Presentation(sys.argv[2])
    print('들여쓰기', fix_hang(_p), '문단 · 단어 끊김 방지', no_word_split(_p), '문단 · 색 정리', fix_colors(_p), '곳')
    _p.save(sys.argv[3])
