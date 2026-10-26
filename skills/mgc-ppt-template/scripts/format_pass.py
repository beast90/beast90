"""format_pass.py — 서식 일괄 적용 부품 (mgc-ppt-template v4.2, 4-C-12)
강조 빨강 #FF0000, 제목 개조식 교체, 표 칸 「- 」 정보 단위(열 판정 include), 빈 칸 「—」 가운데 정렬, 요약 띠 라벨만 굵게,
확인 태그 서술형(tagfix), 순서 번호(number_chain), 바뀐 표만 맞춤(fit_tables).
사용 예: from format_pass import *; ORIG = table_bottoms(p); recolor(slide._element); ...; fit_tables(p, ORIG)"""
import sys, re, copy
sys.path.insert(0, __import__('os').path.dirname(__file__))
from pptx.util import Inches
from pptx.oxml.ns import qn
import pptx_edit as PE
import edit_lib as L
from edit_lib import BULLET, red_text, set_items, cell_w, row_h, wpt

OLD_RED = {'F53148', 'D32F2F', 'E03C31'}
CIRC = '①②③④⑤⑥⑦⑧⑨⑩'
EMPTY = {'—', '-', '–'}
HDR_FILL = 'E2D9CF'
EXC = ('장', '번호', '순서', '단계', '순위', '코드', '카드사', 'Native', 'Webview', 'iOS', 'AOS', '서버', '성격', '필수',
       '다시 부르기', '위치', '쓰는 흐름', '개발 주체', '개발 구분', '상태', '제공 시점', '의미', '담당 영역', '콜백 필드',
       '호출값', '호출 경로', '핏콜라보 서버', 'Android', '고객 화면 문구', '보이는 곳', '근거', '관련 근거', '화면', '결제 화면',
       '요소', 'API', '인증 거래번호', '결제 거래번호', '현행(일반', '앱카드 결제', '내용(호출명)', '보내는 쪽', '주는 쪽',
       '처리 주체', '구분', '값', '항목', '대상', '기능', '방향', '담당', '영역', '구성 요소', '용어', '코드 · 상황', '구현 포인트',
       '사고', '준비 항목', '작업', '개발 과제', '호출값', '보내는 값', '현상', '판정', '주문 결과 상태')

def recolor(root):
    n = 0
    for e in root.iter(qn('a:srgbClr')):
        if e.get('val', '').upper() in OLD_RED: e.set('val', 'FF0000'); n += 1
    return n

def fill(c):
    tcPr = c._tc.find(qn('a:tcPr'))
    f = tcPr.find(qn('a:solidFill')) if tcPr is not None else None
    x = f.find(qn('a:srgbClr')) if f is not None else None
    return x.get('val').upper() if x is not None else ''

def has_header(t): return len(t.rows) > 1 and fill(t.cell(0, 0)) == HDR_FILL
def norm(s): return re.sub(r'\s+', ' ', s).strip()

def include(si, t, ci, col0_ok=()):
    """정보 단위 「- 」를 붙일 열인지: 머리글이 값 · 이름 열이면 제외, 첫 열(이름 열)은 제외"""
    hdr = has_header(t); start = 1 if hdr else 0
    body = [t.cell(ri, ci).text_frame.text for ri in range(start, len(t.rows))]
    if ci == 0 and si not in col0_ok: return False
    if any(b.lstrip().startswith(('• ', '- ')) for b in body): return True
    if hdr:
        h = norm(t.cell(0, ci).text_frame.text)
        if ci > 0 and any(h == x or h.startswith(x + ' ') or h.startswith(x + '(') or (len(x) > 2 and h.startswith(x)) for x in EXC):
            return False
    if any(b.lstrip().startswith(('• ', '- ')) for b in body): return True
    L_ = [len(norm(b)) for b in body if b.strip() and b.strip() not in EMPTY]
    return bool(L_) and sum(L_) / len(L_) >= 10

# ── 문단 → 정보 단위(논리 항목) 복원
def logical(tf):
    reds = red_text(tf); items = []; hasb = False
    for pa, txt in zip(tf.paragraphs, reds):
        s = txt.strip()
        if not s: continue
        head = s.startswith(('• ', '- ', '※'))
        if s.startswith('• '): s = BULLET + s[2:]
        hasb |= head
        if items and not head:
            prev = items[-1]
            plain = s.replace('[[', '').replace(']]', '')
            if ' · ' in prev and ' ' not in re.sub(r'\([^)]*\)', '', plain) and not prev.endswith('·') and not plain.startswith('[') \
                    and prev.count('(') <= prev.count(')') and prev.count('[') <= prev.count(']'):
                items[-1] = prev + ' · ' + s          # 줄 경계에서 뺀 나열 「·」 복원
            else:
                items[-1] = prev + ('' if prev.endswith(('(', '/')) else ' ') + s
        else:
            items.append(s)
    return [re.sub(r'\s+', ' ', i).replace('· ·', '·') for i in items]

TAG = re.compile(r'\[확인필요\(([^):\]]+)\)(?::\s*([^\]]+))?\]')
def tagfix(s, special=None):
    def rep(m):
        who = m.group(1).replace(' · ', ' ↔ ').replace('\n', ' ')
        what = (m.group(2) or '').replace('\n', ' ').strip()
        verb = '검토 필요' if '법무' in who else '확인 필요'
        body = special or f'{who} {what + " " if what else ""}{verb}'
        return f'[[[{body}]]]'
    return TAG.sub(rep, s)

def number_chain(s):
    """「A → B → C」(화살표 2개 이상) 단계에 ①②③ 번호"""
    m = re.match(r'^((?:- )?(?:[^→:]{1,6}: )?)(.*)$', s)
    lead, body = m.group(1), m.group(2)
    if body.count(' → ') < 2 or any(c in body for c in CIRC): return s
    parts = body.split(' → ')
    return lead + ' → '.join(CIRC[k] + x for k, x in enumerate(parts))

def dash_w(size, bold=False): return int(wpt(BULLET, size, bold) * 12700)

def add_dash_run(r0):
    """첫 런 앞에 「- 」: 첫 런이 빨강이면 본문색 런을 따로 넣는다(글머리는 강조하지 않음)"""
    rp = r0._r.find(qn('a:rPr')); c = rp.find('.//' + qn('a:srgbClr')) if rp is not None else None
    if c is None or c.get('val', '').upper() != 'FF0000':
        r0.text = '- ' + r0.text; return
    nr = copy.deepcopy(r0._r); nrp = nr.find(qn('a:rPr'))
    for f in nrp.findall(qn('a:solidFill')): nrp.remove(f)
    sf = nrp.makeelement(qn('a:solidFill'), {}); sf.append(sf.makeelement(qn('a:srgbClr'), {'val': '2E241D'})); nrp.insert(0, sf)
    nr.find(qn('a:t')).text = '- '; r0._r.addprevious(nr)

def prefix_cell(tf):
    """칸을 다시 나누지 않고 「• 」 → 「- 」, 글머리 없는 첫 줄에 「- 」, 이어지는 줄은 「- 」 폭 들여쓰기"""
    paras = [pa for pa in tf.paragraphs if pa.text.strip()]
    if not paras or tf.text.strip() in EMPTY: return False
    changed = False; started = False
    for k, pa in enumerate(paras):
        t = pa.text; r0 = pa.runs[0] if pa.runs else None
        if r0 is None: continue
        size = (r0.font.size.pt if r0.font.size else 11)
        ppr = pa._p.find(qn('a:pPr'))
        if ppr is None: ppr = pa._p.makeelement(qn('a:pPr'), {}); pa._p.insert(0, ppr)
        if t.lstrip().startswith('※'):
            started = False; continue
        if t.startswith('• ') or t.startswith('- '):
            if t.startswith('• '):
                r0.text = '- ' + r0.text[2:] if r0.text.startswith('• ') else r0.text.replace('•', '-', 1)
            for a in ('marL', 'indent'):
                if a in ppr.attrib: del ppr.attrib[a]
            started = True; changed = True
        elif k == 0 or not started:
            add_dash_run(r0)
            for a in ('marL', 'indent'):
                if a in ppr.attrib: del ppr.attrib[a]
            started = True; changed = True
        else:
            ppr.set('marL', str(dash_w(size, bool(r0.font.bold)))); ppr.set('indent', '0'); changed = True
    return changed

def bullets_to_dash(tf):
    """글상자 · 제외 열: 「• 」만 「- 」로 바꾸고 이어지는 줄 들여쓰기 폭을 맞춘다"""
    n = 0; inb = False
    for pa in tf.paragraphs:
        r0 = pa.runs[0] if pa.runs else None
        if r0 is None: continue
        ppr = pa._p.find(qn('a:pPr'))
        if pa.text.startswith('• '):
            r0.text = r0.text.replace('• ', '- ', 1) if r0.text.startswith('• ') else r0.text.replace('•', '-', 1); n += 1; inb = True
        elif pa.text.startswith(('- ', '※')) or not pa.text.strip(): inb = pa.text.startswith('- ')
        elif inb and ppr is not None and ppr.get('marL') not in (None, '0'):
            size = r0.font.size.pt if r0.font.size else 11
            ppr.set('marL', str(dash_w(size, bool(r0.font.bold))))
    return n

def center_empty(t):
    n = 0
    for row in t.rows:
        for c in row.cells:
            if c.text_frame.text.strip() in EMPTY:
                for pa in c.text_frame.paragraphs:
                    ppr = pa._p.find(qn('a:pPr'))
                    if ppr is None: ppr = pa._p.makeelement(qn('a:pPr'), {}); pa._p.insert(0, ppr)
                    ppr.set('algn', 'ctr')
                    for a in ('marL', 'indent'):
                        if a in ppr.attrib: del ppr.attrib[a]
                n += 1
    return n

def strip_label_bold(sh):
    """요약 띠 「라벨   내용」: 라벨만 굵게, 내용은 보통(사용자 p.4 기준)"""
    tf = sh.text_frame; pa = tf.paragraphs[0]; t = pa.text
    m = re.match(r'^(.+?)(   )(.+)$', t)
    if not m or not pa.runs: return False
    tmpl = copy.deepcopy(pa.runs[0]._r)
    for r in list(pa._p.findall(qn('a:r'))): pa._p.remove(r)
    end = pa._p.find(qn('a:endParaRPr'))
    for txt, b in ((m.group(1), '1'), (m.group(2), '1'), (m.group(3), '0')):
        nr = copy.deepcopy(tmpl); nr.find(qn('a:t')).text = txt; nr.find(qn('a:rPr')).set('b', b)
        (end.addprevious(nr) if end is not None else pa._p.append(nr))
    return True

def set_runs_text(tf, text):
    """제목처럼 한 문단 글: 첫 런 서식으로 글만 바꾼다"""
    pa = tf.paragraphs[0]; r0 = pa.runs[0]; r0.text = text
    rp = r0._r.find(qn('a:rPr'))
    for f in rp.findall(qn('a:solidFill')): rp.remove(f)
    sf = rp.makeelement(qn('a:solidFill'), {}); sf.append(sf.makeelement(qn('a:srgbClr'), {'val': '2E241D'})); rp.insert(0, sf)
    for r in pa.runs[1:]: r._r.getparent().remove(r._r)
    for extra in tf.paragraphs[1:]: extra._p.getparent().remove(extra._p)

def width_over(tf, w_in):
    avail = w_in * 72 * 0.95
    for pa in tf.paragraphs:
        ppr = pa._p.find(qn('a:pPr')); ml = int(ppr.get('marL', '0')) / 12700 if ppr is not None else 0
        w = 0
        for r in pa.runs:
            sz = r.font.size.pt if r.font.size else 11; b = bool(r.font.bold)
            w += wpt(r.text, sz, b)
        if w + ml > avail: return True
    return False

def table_bottoms(p):
    return {(si, sh.shape_id): ((sh.top + sum(r.height for r in sh.table.rows)) / 914400, ttext(sh.table)) for si, s in enumerate(p.slides, 1) for sh in s.shapes if sh.has_table}

def ttext(t): return '|'.join(c.text_frame.text for r in t.rows for c in r.cells)

def fit_tables(p, orig=None):
    def set_tbl_size(t, size):
        for e in list(t._tbl.iter(qn('a:rPr'))) + list(t._tbl.iter(qn('a:endParaRPr'))): e.set('sz', str(int(size * 100)))
    out = []
    for si, s in enumerate(p.slides, 1):
        for sh in [x for x in s.shapes if x.has_table]:
            t = sh.table; top = sh.top / 914400; Lx, R = sh.left, sh.left + sh.width
            if orig and (si, sh.shape_id) in orig and orig[(si, sh.shape_id)][1] == ttext(t): continue
            for ri in range(len(t.rows)):
                h = Inches(row_h(t, ri, 0.30))
                if h > t.rows[ri].height: t.rows[ri].height = h
            below = [o.top / 914400 for o in s.shapes if o is not sh and o.top / 914400 > top + 0.3
                     and not (o.left + o.width <= Lx or R <= o.left) and not (o.has_text_frame and re.match(r'^\d+ / \d+$', o.text_frame.text.strip()))]
            limit = min(below + [6.98]) - 0.06
            if orig and (si, sh.shape_id) in orig:
                ob, otx = orig[(si, sh.shape_id)]
                if otx == ttext(t): continue
                limit = max(limit, ob)
            def bot(): return top + sum(r.height for r in t.rows) / 914400
            if bot() <= limit + 0.005: continue
            for ri in range(len(t.rows)): t.rows[ri].height = Inches(row_h(t, ri, 0.30))
            size, _ = PE.run_props(t.cell(len(t.rows) - 1, len(t.columns) - 1).text_frame)
            while bot() > limit + 0.005 and size > 10.5:
                size -= 0.5; set_tbl_size(t, size)
                for ri, row in enumerate(t.rows):
                    for ci, c in enumerate(row.cells):
                        if c.text_frame.text.strip() and width_over(c.text_frame, cell_w(t, ci)):
                            set_items(c.text_frame, logical(c.text_frame), cell_w(t, ci))
                    t.rows[ri].height = Inches(row_h(t, ri, 0.30))
            out.append((si, size, round(bot(), 2), round(limit, 2)))
    return out
