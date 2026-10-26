"""move_remap.py — 장 하나를 다른 위치로 옮기고 쪽 번호 · 간지 목록 · 본문 「N장」 · 「a~b」 범위 · 노트 · 목차 범위를 새 번호로 바꾼다 (v4.1)
사용: python3 move_remap.py in.pptx out.pptx 옮길장 새위치   (예: 78 5)
규격서 · 정의서 같은 다른 문서의 장 번호(「인증 모듈 2.1.2 6장」 「정의서 v0.2 12장」)는 바꾸지 않는다(DOC 제외 규칙)"""
import sys, re, copy
sys.path.insert(0, __import__('os').path.dirname(__file__))
from pptx import Presentation
from pptx.util import Inches
import pptx_edit as PE

SRC, OUT = sys.argv[1], sys.argv[2]
OLD_POS, NEW_POS = int(sys.argv[3]), int(sys.argv[4])
def f(n):
    if n == OLD_POS: return NEW_POS
    if NEW_POS <= n < OLD_POS: return n + 1
    if OLD_POS < n <= NEW_POS: return n - 1
    return n

p = Presentation(SRC)
ids = p.slides._sldIdLst; lst = list(ids)
el = lst[OLD_POS - 1]; ids.remove(el); ids.insert(NEW_POS - 1, el)
N = len(p.slides)

PG = re.compile(r'^\d+ / \d+$')
LIST = re.compile(r'^(\d+)(\s+)(.*)$')
REF = re.compile(r'(?<![\d.~])(\d{1,2})((?:\s*·\s*\d{1,2})*)\s*장')
RNG = re.compile(r'(?<![\d.])(\d{1,2})~(\d{1,2})(?=\s*장|\))')
DOC = re.compile(r'(규격|정의서|모듈|간편결제|로그인|가이드|요약서|답변|v\d|\d\.\d|안내|화면설계|Flipbook|부록|별표)[^장]{0,6}$')
def is_doc(t, start): return bool(DOC.search(t[max(0, start - 14):start]))
def remap_text(t):
    t = RNG.sub(lambda m: m.group(0) if is_doc(t, m.start()) else f'{f(int(m.group(1)))}~{f(int(m.group(2)))}', t)
    def one(m):
        if is_doc(t, m.start()): return m.group(0)
        nums = [m.group(1)] + re.findall(r'\d{1,2}', m.group(2))
        seps = re.findall(r'\s*·\s*', m.group(2))
        out = str(f(int(nums[0])))
        for s, n in zip(seps, nums[1:]): out += s + str(f(int(n)))
        return out + m.group(0)[len(m.group(1) + m.group(2)):]
    return REF.sub(one, t)

def runs_all(s):
    for sh in s.shapes:
        if sh.has_text_frame:
            for pa in sh.text_frame.paragraphs:
                for r in pa.runs: yield sh, r
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    for pa in c.text_frame.paragraphs:
                        for r in pa.runs: yield sh, r

for i, s in enumerate(p.slides, 1):
    for sh, r in runs_all(s):
        t = r.text
        if PG.match(t.strip()): r.text = f'{i} / {N}'; continue
        if sh.name == '간지 목록':
            m = LIST.match(t)
            if m: r.text = f'{f(int(m.group(1)))}{m.group(2)}{m.group(3)}'
            continue
        t2 = remap_text(t)
        if t2 != t: r.text = t2
    if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
        tf = s.notes_slide.notes_text_frame
        for pa in tf.paragraphs:
            for r in pa.runs:
                t2 = remap_text(r.text)
                if t2 != r.text: r.text = t2

# 2장 목차의 「장」 열(장 글자 없음, 목차가 2장에 있을 때)
t2 = PE.tbl(p.slides[1]).table
for ri in range(1, len(t2.rows)):
    c = t2.cell(ri, 1); m = re.match(r'^(\d+)~(\d+)$', c.text.strip())
    if m: PE.set_cell(t2, ri, 1, f'{f(int(m.group(1)))}~{f(int(m.group(2)))}')
p.save(OUT)
print('moved', N)
