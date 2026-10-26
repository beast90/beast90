"""장식·화면 밖 요소 검사. 사용: python3 check_clean.py out.pptx
- 슬라이드 밖으로 나가거나 걸친 도형
- 순번 원(0.4in)보다 큰 원형 도형(장식 원)
- 기능 보고형에서 쓰지 않는 장식용 색(4A382C) 도형
문제 0건이면 '깨끗함' 출력"""
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
E = 914400
p = Presentation(sys.argv[1]); W, H = p.slide_width / E, p.slide_height / E
bad = 0
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        x, y, w, h = (v / E for v in (sh.left, sh.top, sh.width, sh.height))
        msg = []
        if x < -0.001 or y < -0.001 or x + w > W + 0.001 or y + h > H + 0.001: msg.append('화면 밖')
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                if sh.auto_shape_type == MSO_SHAPE.OVAL and max(w, h) > 0.45: msg.append('장식 원')
            except Exception: pass
            try:
                if sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == '4A382C': msg.append('장식 색')
            except Exception: pass
        if msg:
            bad += 1; print(f'{i}장 {sh.name} ({x:.2f},{y:.2f},{w:.2f},{h:.2f}): {", ".join(msg)}')
print('깨끗함' if bad == 0 else f'문제 {bad}건')
