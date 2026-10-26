"""tidy_tables.py — 이미 만든 덱의 표를 「깔끔한 표」 기준으로 정리 (mgc-ppt-template v5.5, references/text-fit.md 「깔끔한 표 · 상자」)
사용: python3 tidy_tables.py 입력.pptx 출력.pptx      또는  from tidy_tables import run; run(p)
- 한 줄에 들어가는데 두 줄 이상으로 나뉜 「- 」 정보를 한 줄로 합친다(줄바꿈 폭 보정 pptx_edit.SLACK)
- 열 폭 합은 그대로 두고 문장 열 폭을 다시 나눠 더 많은 정보가 한 줄에 들어가면 그 폭을 쓴다(balance_cols)
- 줄 수에 맞춰 행 높이를 다시 잡는다. 표 하단은 원래보다 내려가지 않는다(내려가면 그 표는 되돌림)
- 건드리지 않는 것: 간지, 첫 열, 「- 」로 시작하지 않는 칸(「번호 / 이름 / 호출명」처럼 일부러 나눈 칸), 서식이 섞인 칸(빨강 제외)"""
import sys, re, copy
sys.path.insert(0, __import__('os').path.dirname(__file__))
from pptx.util import Inches
from pptx.oxml.ns import qn
import pptx_edit as PE

def _has_header(t):
    f = t.cell(0, 0)._tc.find('.//' + qn('a:srgbClr'))
    return f is not None and f.get('val', '').upper() == 'E2D9CF'   # 머리행(doc-surface-tan)

def _bullet_cell(c):
    ps = [pa for pa in c.text_frame.paragraphs if pa.text.strip()]
    return bool(ps) and ps[0].text.startswith('- ')

def _uniform(c):
    keys = set()
    for pa in c.text_frame.paragraphs:
        for r in pa.runs:
            if PE._isred(r._r) or not r.text.strip(): continue
            keys.add((r.font.bold, r.font.size, r.font.italic))
    return len(keys) <= 1

def _lines_needed(c, w_in):
    size, bold = PE.run_props(c.text_frame)
    return sum(len(PE.wrap(re.sub(r'\[\[|\]\]', '', it), w_in, size, bold)) for it in PE.merged_text(c).split('\n') if it.strip())

def _paras(c): return len([pa for pa in c.text_frame.paragraphs if pa.text.strip()])

def _col_w(t, ci): return t.columns[ci].width / 914400 - 0.24

def _fits_all(t, ci):
    """칸 폭을 바꿨을 때 「- 」가 아닌 칸의 기존 줄이 넘치지 않는지"""
    for ri in range(len(t.rows)):
        c = t.cell(ri, ci)
        if _bullet_cell(c): continue
        size, bold = PE.run_props(c.text_frame)
        if any(PE.wpt(pa.text, size, bold) > _col_w(t, ci) * 72 * PE.SLACK for pa in c.text_frame.paragraphs): return False
    return True

def tidy_table(sh, minh=0.366):
    t = sh.table; nr, nc = len(t.rows), len(t.columns)
    if nr < 2: return None
    r0 = 1 if _has_header(t) else 0
    keep = copy.deepcopy(sh._element); bot0 = PE.bottom(sh)
    before = sum(_paras(t.cell(ri, ci)) for ri in range(nr) for ci in range(nc))
    w0 = [c.width for c in t.columns]
    # 1) 열 폭 다시 나누기: 「- 」 칸 줄 수가 줄고 다른 칸이 넘치지 않을 때만
    def total_lines():
        return sum(_lines_needed(t.cell(ri, ci), _col_w(t, ci)) for ri in range(r0, nr) for ci in range(1, nc)
                   if _bullet_cell(t.cell(ri, ci)) and _uniform(t.cell(ri, ci)))
    base = total_lines()
    if PE.balance_cols(t) is not None:
        changed = [ci for ci in range(nc) if t.columns[ci].width != w0[ci]]
        if total_lines() >= base or not all(_fits_all(t, ci) for ci in changed):
            for ci in range(nc): t.columns[ci].width = w0[ci]
    changed = {ci for ci in range(nc) if t.columns[ci].width != w0[ci]}
    # 2) 「- 」 칸 다시 줄바꿈
    touched = set()
    for ri in range(r0, nr):
        for ci in range(1, nc):
            c = t.cell(ri, ci)
            if not _bullet_cell(c) or not _uniform(c): continue
            need = _lines_needed(c, _col_w(t, ci))
            if ci in changed or need < _paras(c):
                PE.set_cell(t, ri, ci, PE.merged_text(c)); touched.add(ri)
    # 3) 행 높이: 줄 수에 맞춤(빈 줄 공간 없앰)
    for ri in range(r0, nr):
        row = t.rows[ri]
        size = max(PE.run_props(c.text_frame)[0] for c in row.cells)
        need = max(_paras(c) * PE.run_props(c.text_frame)[0] * 1.3 / 72 for c in row.cells) + 0.14
        if ri in touched or row.height / 914400 - need > 0.8 * size * 1.3 / 72:
            row.height = Inches(max(minh, need))
    after = sum(_paras(t.cell(ri, ci)) for ri in range(nr) for ci in range(nc))
    if PE.bottom(sh) > bot0 + 0.005:   # 하단이 내려가면 되돌림
        sh._element.getparent().replace(sh._element, keep); return ('되돌림', before, before)
    return ('정리', before, after) if (after != before or changed or PE.bottom(sh) < bot0 - 0.005) else None

def run(p):
    out = []
    for si, s in enumerate(p.slides, 1):
        if any(x.name == '간지 목록' for x in s.shapes): continue
        for sh in [x for x in s.shapes if x.has_table]:
            r = tidy_table(sh)
            if r: out.append((si,) + r)
    return out

if __name__ == '__main__':
    from pptx import Presentation
    p = Presentation(sys.argv[1]); res = run(p); p.save(sys.argv[2])
    for r in res: print(f'{r[0]}장 {r[1]}: 문단 {r[2]} → {r[3]}')
    print(f'정리한 표 {sum(1 for r in res if r[1] == "정리")}개, 되돌린 표 {sum(1 for r in res if r[1] == "되돌림")}개')
